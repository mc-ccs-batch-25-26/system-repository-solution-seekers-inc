import io
from django.http import HttpResponse
from django.utils.timezone import localtime
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from apps.users.permissions import IsAdminOrOfficial
from apps.beneficiaries.models import Beneficiary, Household, Family
from apps.cycles.models import ProgramCycle, CycleApplication, ParticipationRecord, DailyParticipationRecord
from apps.audit.models import AuditLog
from apps.criteria.models import Criterion
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color='1a56db', end_color='1a56db', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
FORMULA_PREFIXES = ('=', '+', '-', '@')


def _style_headers(ws, row_num):
    for cell in ws[row_num]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _init_sheet(ws, title, headers):
    """Write a bold title in row 1 (merged), a blank row 2, and styled column headers in row 3."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).value = title
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(headers)
    _style_headers(ws, 3)


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 55)


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _append_safe(ws, values):
    ws.append([_safe_cell(value) for value in values])


def _safe_filename(value):
    return ''.join(ch if ch.isalnum() or ch in {'-', '_', '.'} else '_' for ch in value)


def _excel_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = f'attachment; filename="{_safe_filename(filename)}"'
    return resp


class CycleRankingReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOfficial]

    def get(self, request):
        cycle_id = request.query_params.get('cycle_id')
        if not cycle_id:
            return Response({'detail': 'cycle_id is required.'}, status=400)

        try:
            cycle = ProgramCycle.objects.get(pk=cycle_id)
        except (ProgramCycle.DoesNotExist, ValueError):
            return Response({'detail': 'Cycle not found.'}, status=404)

        applications = (
            CycleApplication.objects.filter(cycle=cycle)
            .select_related('beneficiary')
            .prefetch_related('beneficiary__indicators__criterion')
            .order_by('rank_position', '-computed_score')
        )
        criteria = list(Criterion.objects.filter(is_active=True).order_by('name'))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ranking'

        total_cols = max(5 + len(criteria), 5)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        ws.cell(1, 1).value = f'Cycle Ranking Report - {cycle.cycle_name}'
        ws.cell(1, 1).font = Font(bold=True, size=12)
        ws.cell(2, 1).value = (
            f'Slots: {cycle.slots}  |  Application: {cycle.start_date} to {cycle.end_date}'
            f'  |  Work: {cycle.work_start_date or ""} to {cycle.work_end_date or ""}'
        )
        ws.cell(2, 1).font = Font(italic=True, size=9)

        headers = ['Rank', 'Beneficiary Name', 'Score', 'Status', 'Application Date'] + [c.name for c in criteria]
        ws.append([])
        ws.append(headers)
        _style_headers(ws, 4)

        for app in applications:
            indicator_map = {
                str(ind.criterion_id): round(float(ind.value), 4)
                for ind in app.beneficiary.indicators.all()
            }
            _append_safe(ws, [
                app.rank_position or '',
                app.beneficiary.full_name,
                round(float(app.computed_score), 4) if app.computed_score else '',
                app.get_status_display(),
                str(app.application_date),
            ] + [indicator_map.get(str(c.id), '') for c in criteria])

        _auto_width(ws)
        safe_name = _safe_filename(cycle.cycle_name.replace(' ', '_'))
        return _excel_response(wb, f'cycle_ranking_{safe_name}.xlsx')


class BeneficiaryMasterlistReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOfficial]

    def get(self, request):
        beneficiaries = (
            Beneficiary.objects
            .select_related('family__household')
            .order_by('full_name')
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Beneficiary Masterlist'
        _init_sheet(ws, 'Beneficiary Masterlist', [
            'Full Name', 'Age', 'Gender', 'Civil Status', 'Role in Household',
            'Sectors', 'Employment Status', 'Monthly Income',
            'Lot Area (sqm)', 'Dependents', 'Housing Condition',
            'Structure Condition', 'Tenure Status', 'Electrical Condition', 'Water Source',
            'TUPAD Eligible', 'Household Code', 'Purok / Address',
        ])

        for b in beneficiaries:
            household = b.family.household if b.family_id else None
            _append_safe(ws, [
                b.full_name,
                b.age,
                b.get_gender_display(),
                b.get_civil_status_display(),
                b.get_role_display(),
                ', '.join(b.sectors) if b.sectors else '',
                b.get_employment_status_display(),
                float(b.monthly_income),
                float(b.lot_area_sqm) if b.lot_area_sqm is not None else '',
                b.num_dependents,
                b.get_housing_condition_display(),
                b.get_structure_condition_display() if b.structure_condition else '',
                b.get_tenure_status_display() if b.tenure_status else '',
                b.get_electrical_condition_display() if b.electrical_condition else '',
                b.get_water_source_display() if b.water_source else '',
                'Yes' if b.is_tupad_eligible else 'No',
                household.household_code if household else '',
                ((household.purok + ' - ' if household and household.purok else '') + (household.address if household else '')).strip(),
            ])

        _auto_width(ws)
        return _excel_response(wb, 'beneficiary_masterlist.xlsx')


class ParticipationHistoryReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOfficial]

    def get(self, request):
        cycle_id = request.query_params.get('cycle_id')
        qs = (
            ParticipationRecord.objects
            .select_related('beneficiary', 'cycle', 'recorded_by')
            .order_by('cycle__start_date', 'beneficiary__full_name')
        )
        if cycle_id:
            qs = qs.filter(cycle_id=cycle_id)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Participation History'
        _init_sheet(ws, 'Participation History Report', [
            'Beneficiary Name', 'Cycle', 'Project', 'Days Worked', 'Start Date', 'End Date', 'Recorded By',
        ])

        for rec in qs:
            recorder = rec.recorded_by
            recorder_name = getattr(recorder, 'full_name', None) or str(recorder)
            _append_safe(ws, [
                rec.beneficiary.full_name,
                rec.cycle.cycle_name,
                rec.project_name,
                rec.days_worked,
                str(rec.participation_start),
                str(rec.participation_end),
                recorder_name,
            ])

        ws_daily = wb.create_sheet('Daily Records')
        _init_sheet(ws_daily, 'Daily Participation Records', [
            'Beneficiary Name', 'Cycle', 'Project', 'Work Date', 'Status', 'Hours Worked', 'Photo', 'Remarks', 'Recorded By',
        ])
        daily_qs = (
            DailyParticipationRecord.objects
            .select_related('participation__beneficiary', 'participation__cycle', 'recorded_by')
            .order_by('work_date', 'participation__beneficiary__full_name')
        )
        if cycle_id:
            daily_qs = daily_qs.filter(participation__cycle_id=cycle_id)
        for row in daily_qs:
            recorder = row.recorded_by
            recorder_name = getattr(recorder, 'full_name', None) or str(recorder)
            _append_safe(ws_daily, [
                row.participation.beneficiary.full_name,
                row.participation.cycle.cycle_name,
                row.participation.project_name,
                str(row.work_date),
                row.get_status_display(),
                float(row.hours_worked),
                row.photo.url if row.photo else '',
                row.remarks,
                recorder_name,
            ])

        _auto_width(ws)
        _auto_width(ws_daily)
        return _excel_response(wb, 'participation_history.xlsx')


class AuditTrailReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOfficial]

    def get(self, request):
        start = request.query_params.get('start_date')
        end = request.query_params.get('end_date')

        qs = AuditLog.objects.select_related('user').order_by('-timestamp')
        if start:
            qs = qs.filter(timestamp__date__gte=start)
        if end:
            qs = qs.filter(timestamp__date__lte=end)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Audit Trail'
        _init_sheet(ws, 'Audit Trail Report', [
            'Timestamp', 'User', 'Action', 'Target Table', 'Target ID', 'Details',
        ])

        for log in qs:
            _append_safe(ws, [
                localtime(log.timestamp).strftime('%Y-%m-%d %H:%M:%S'),
                getattr(log.user, 'full_name', None) or str(log.user) if log.user else 'System',
                log.action,
                log.target_table,
                log.target_id or '',
                str(log.details) if log.details else '',
            ])

        _auto_width(ws)
        return _excel_response(wb, 'audit_trail.xlsx')


class HouseholdReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOfficial]

    def get(self, request):
        households = (
            Household.objects
            .prefetch_related('families__members')
            .order_by('household_code')
        )

        wb = openpyxl.Workbook()

        # Sheet 1 - Households
        ws1 = wb.active
        ws1.title = 'Households'
        _init_sheet(ws1, 'Household List', [
            'Household Code', 'House/Block/Lot No.', 'Street Name', 'Purok/Sitio',
            'Landmark', 'Length of Residency', 'Full Address', 'Status',
        ])
        for h in households:
            _append_safe(ws1, [
                h.household_code,
                h.house_no_or_blk_lot_no,
                h.street_name,
                h.purok_sitio,
                h.landmark,
                h.length_of_residency if h.length_of_residency is not None else '',
                h.address,
                h.get_status_display(),
            ])
        _auto_width(ws1)

        # Sheet 2 - Families
        ws2 = wb.create_sheet('Families')
        _init_sheet(ws2, 'Family List', ['Household Code', 'Family No.', 'Monthly Income Bracket'])
        for h in households:
            for f in h.families.all():
                _append_safe(ws2, [h.household_code, f.family_number, f.get_monthly_income_bracket_display()])
        _auto_width(ws2)

        # Sheet 3 - Members
        ws3 = wb.create_sheet('Members')
        _init_sheet(ws3, 'Household Members', [
            'Household Code', 'Family No.', 'Full Name', 'Role', 'Age', 'Gender',
            'Civil Status', 'Sectors', 'Employment Status', 'Monthly Income',
            'Lot Area (sqm)', 'Dependents', 'Housing Condition', 'Structure Condition',
            'Tenure Status', 'Electrical Condition', 'Water Source', 'TUPAD Eligible', 'Contact No.',
        ])
        for h in households:
            for f in h.families.all():
                for m in f.members.all():
                    _append_safe(ws3, [
                        h.household_code,
                        f.family_number,
                        m.full_name,
                        m.get_role_display(),
                        m.age,
                        m.get_gender_display(),
                        m.get_civil_status_display(),
                        ', '.join(m.sectors) if m.sectors else '',
                        m.get_employment_status_display(),
                        float(m.monthly_income),
                        float(m.lot_area_sqm) if m.lot_area_sqm is not None else '',
                        m.num_dependents,
                        m.get_housing_condition_display(),
                        m.get_structure_condition_display() if m.structure_condition else '',
                        m.get_tenure_status_display() if m.tenure_status else '',
                        m.get_electrical_condition_display() if m.electrical_condition else '',
                        m.get_water_source_display() if m.water_source else '',
                        'Yes' if m.is_tupad_eligible else 'No',
                        m.contact_number,
                    ])
        _auto_width(ws3)

        return _excel_response(wb, 'household_report.xlsx')
